from datetime import date, time
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import patch

import pytest
from flask import Flask, Blueprint
from flask_login import LoginManager
from openpyxl import load_workbook
from app.extensions import db
from app.models import Usuario, Prefeitura, Solicitacao, Equipe, OrdemServico
from app.modules.chatbot import exports
from app.modules.chatbot.routes import register_routes
from app.modules.consultas import service


@pytest.fixture
def app(tmp_path):
    app = Flask(__name__, instance_path=str(tmp_path))
    app.config.update(TESTING=True, SECRET_KEY='only-test', SQLALCHEMY_DATABASE_URI='sqlite:///:memory:')
    db.init_app(app)
    manager = LoginManager(app)
    manager.user_loader(lambda uid: db.session.get(Usuario, int(uid)))
    bp = Blueprint('main', __name__)
    register_routes(bp)
    app.register_blueprint(bp)
    with app.app_context():
        db.create_all()
        db.session.add(Prefeitura(id=1,nome='Prefeitura de teste',slug='teste'))
        for uid, role in [(1,'dev'), (2,'diretor'), (3,'admin'), (4,'piloto_agro')]:
            db.session.add(Usuario(id=uid, login=f'user{uid}', nome_uvis='Teste', senha_hash='unused',tipo_usuario=role,prefeitura_id=1))
        db.session.add(Equipe(id=1,nome_equipe='Teste',prefeitura_id=1,regiao='NORTE'))
        db.session.flush()
        row = Solicitacao(id=1,usuario_id=1,prefeitura_id=1,data_agendamento=date(2026,10,17),hora_agendamento=time(10),
            foco='Piscina',tipo_visita='Aedes',cep='01000000',logradouro='Rua Teste',bairro='Centro',cidade='Cidade',uf='SP',status='PENDENTE')
        db.session.add(row)
        db.session.flush()
        db.session.add(OrdemServico(solicitacao_id=1,equipe_id=1,data_aplicacao=date(2026,10,20),situacao_aplicacao='CONCLUIDA',identificador_os='OS-TESTE'))
        db.session.commit()
    yield app
    with app.app_context():
        db.session.remove(); db.drop_all(); db.engine.dispose()


def client(app, uid=1):
    c = app.test_client()
    with c.session_transaction() as session:
        session['_user_id'] = str(uid); session['_fresh'] = True
    return c


def plan(**changes):
    return dict({'source':'solicitacoes','format':'excel','inicio':'2026-10-01','fim':'2026-10-31','bairro':'','status':'PENDENTE'}, **changes)


def prepare(app, uid=1, **changes):
    with app.test_request_context():
        return exports.prepare(db.session.get(Usuario,uid),plan(**changes))


@pytest.mark.parametrize("uid", [1, 2])
def test_excel_requires_confirmation_and_matches_preview(app, uid):
    preview = prepare(app, uid=uid)
    assert preview['total'] == 1
    assert preview['filters']['Status da solicitação'] == 'PENDENTE'
    c = client(app, uid)
    assert c.post('/api/azul/exportar',json={'token':preview['token']}).status_code == 400
    response = c.post('/api/azul/exportar',json={'token':preview['token'],'confirm':True})
    assert response.status_code == 200
    assert response.headers['Cache-Control'] == 'private, no-store'
    book = load_workbook(BytesIO(response.data))
    assert book['Consulta'].max_row == 2
    assert 'PENDENTE' in [cell.value for cell in book['Consulta'][2]]


def test_pdf_and_execution_date_filters(app):
    preview = prepare(app, source='ordens_servico',format='pdf',inicio='2026-10-20',fim='2026-10-20',status='')
    assert preview['total'] == 1
    assert preview['filters']['Data utilizada'] == 'Aplicação da OS'
    with patch('app.modules.consultas.exporters.report_prefeitura_id', return_value=None):
        response = client(app).post('/api/azul/exportar',json={'token':preview['token'],'confirm':True})
    assert response.status_code == 200
    assert response.data.startswith(b'%PDF')


@pytest.mark.parametrize('uid',[3,4])
def test_other_roles_cannot_generate_even_with_valid_token(app,uid):
    preview=prepare(app)
    assert client(app,uid).post('/api/azul/exportar',json={'token':preview['token'],'confirm':True}).status_code == 403
    with app.app_context():
        assert exports.export_tool(db.session.get(Usuario,uid)) == []


def test_token_is_user_bound_and_tamper_protected(app):
    preview=prepare(app)
    assert client(app,2).post('/api/azul/exportar',json={'token':preview['token'],'confirm':True}).status_code == 403
    assert client(app).post('/api/azul/exportar',json={'token':preview['token']+'x','confirm':True}).status_code == 400
    with patch('itsdangerous.timed.TimestampSigner.get_timestamp',return_value=1):
        expired=prepare(app)
    assert client(app).post('/api/azul/exportar',json={'token':expired['token'],'confirm':True}).status_code == 400


def test_changed_data_requires_new_preview(app):
    preview=prepare(app)
    with app.app_context():
        db.session.get(Solicitacao,1).bairro='Novo bairro';db.session.commit()
    assert client(app).post('/api/azul/exportar',json={'token':preview['token'],'confirm':True}).status_code == 409


def test_revoked_permission_blocks_preview_and_download(app):
    preview=prepare(app)
    with patch.object(exports,'can_access_endpoint',return_value=False):
        assert client(app).post('/api/azul/exportar',json={'token':preview['token'],'confirm':True}).status_code == 403
        with app.test_request_context():
            assert exports.allowed_sources(db.session.get(Usuario,1)) == []


@pytest.mark.parametrize('changes', [{'source':'agro'}, {'format':'sql'}, {'inicio':'2026-99-99'}, {'inicio':'2026-11-01'}, {'status':[]}, {'sql':'DROP TABLE usuarios'}])
def test_invalid_plans_are_rejected(app,changes):
    with app.test_request_context(), pytest.raises(service.QueryError):
        exports.prepare(db.session.get(Usuario,1),plan(**changes))


def test_empty_result_is_not_offered_for_download(app):
    with app.test_request_context(), pytest.raises(service.QueryError,match='Nenhum registro'):
        exports.prepare(db.session.get(Usuario,1),plan(bairro='Inexistente'))


def test_real_query_is_scoped_by_existing_service(app):
    with app.test_request_context():
        user=db.session.get(Usuario,2)
        with patch.object(service,'base_query', wraps=service.base_query) as scoped:
            result=exports.prepare(user,plan())
            assert result['total']==1
            assert scoped.call_args.args[0] is user


def test_model_only_prepares_and_never_generates_file(app):
    app.config['OPENAI_API_KEY']='test'
    import json
    from app.modules.chatbot import azul
    with patch.object(azul,'OpenAI') as sdk, patch.object(exports.exporters,'excel') as excel:
        sdk.return_value.__enter__.return_value.responses.create.return_value=SimpleNamespace(
            status='completed',output_text='',output=[SimpleNamespace(type='function_call',name='preparar_exportacao',arguments=json.dumps(plan()))])
        response=client(app).post('/api/azul/ajuda',json={'message':'Gere Excel de solicitações pendentes de outubro de 2026'})
        assert response.status_code==200
        assert response.json['export_preview']['total']==1
        excel.assert_not_called()


@pytest.mark.parametrize('uid', [1, 2, 3])
def test_report_start_asks_filters_without_model_or_query(app, uid):
    from app.modules.chatbot import azul
    with patch.object(azul, 'OpenAI') as sdk, patch.object(exports, 'prepare') as prepare_report:
        response = client(app, uid).post('/api/azul/ajuda', json={
            'message': 'Quero gerar um relatório.', 'intent': 'start_report'})
        assert response.status_code == 200
        assert response.json['export_preview'] is None
        if uid in (1, 2):
            for text in ('Solicitações municipais', 'Ordens de serviço', 'Excel ou PDF', 'Qual período', 'bairro', 'status'):
                assert text in response.json['answer']
        else:
            assert 'não está disponível' in response.json['answer']
        sdk.assert_not_called()
        prepare_report.assert_not_called()


def test_report_start_lists_only_authorized_sources(app):
    with patch.object(exports, 'allowed_sources', return_value=['solicitacoes']):
        response = client(app).post('/api/azul/ajuda', json={
            'message': 'Quero gerar um relatório.', 'intent': 'start_report'})
    assert '• Solicitações municipais' in response.json['answer']
    assert '• Ordens de serviço' not in response.json['answer']


def test_repeated_downloads_are_limited_before_query(app):
    preview = prepare(app)
    c = client(app)
    for _ in range(3):
        assert c.post('/api/azul/exportar', json={'token': preview['token'], 'confirm': True}).status_code == 200
    with patch.object(service, 'execute') as query:
        response = c.post('/api/azul/exportar', json={'token': preview['token'], 'confirm': True})
        assert response.status_code == 429
        query.assert_not_called()


def test_download_fails_closed_when_quota_storage_unavailable(app):
    import sqlite3
    from app.modules.chatbot import azul
    preview = prepare(app)
    with patch.object(azul, 'consume_help_quota', side_effect=sqlite3.OperationalError('private details')), patch.object(service, 'execute') as query:
        response = client(app).post('/api/azul/exportar', json={'token': preview['token'], 'confirm': True})
    assert response.status_code == 503
    assert 'private details' not in response.get_data(as_text=True)
    query.assert_not_called()


def test_download_csrf_rejects_missing_or_foreign_token(app):
    from app.shared.csrf_security import register_csrf_security, CSRF_SESSION_KEY
    app.config.update(SECRET_KEY='test-only-secret-key-with-at-least-32-characters', CSRF_PROTECTION_ENABLED=True)
    register_csrf_security(app)
    preview = prepare(app)
    c = client(app)
    with c.session_transaction() as session:
        session[CSRF_SESSION_KEY] = 'session-specific-token'
    payload = {'token': preview['token'], 'confirm': True}
    with patch.object(service, 'execute') as query:
        assert c.post('/api/azul/exportar', json=payload).status_code == 403
        assert c.post('/api/azul/exportar', json=payload, headers={'X-CSRFToken': 'foreign-token'}).status_code == 403
        query.assert_not_called()
    assert c.post('/api/azul/exportar', json=payload, headers={'X-CSRFToken': 'session-specific-token'}).status_code == 200


def test_export_daily_quota_is_separate_and_per_user(app):
    from app.modules.chatbot import azul
    with app.app_context():
        for minute in range(50):
            with patch.object(azul.time, 'time', return_value=86400 + minute * 60):
                assert azul.consume_help_quota(1, export=True)
        with patch.object(azul.time, 'time', return_value=86400 + 51 * 60):
            assert not azul.consume_help_quota(1, export=True)
            assert azul.consume_help_quota(2, export=True)
            assert azul.consume_help_quota(1)
