"""Migração de dados legados (TR 4.7.1) em três etapas rastreáveis.

I   Extração: CSV/TXT, XLSX ou DBF do sistema legado → texto CSV UTF-8.
II  Validação: mapeamento de colunas, padronização, correções automáticas,
    duplicados, obrigatórios, valores inválidos e enriquecimento.
III Adequação: layout interno conferido pelos validadores já existentes
    (preview_csv, sinan_import.parse_cases, planning.read_csv) e carga pelos
    mesmos stores das importações de vigilância.
"""

import csv
import hashlib
import io
import math
import os
import re
import struct
import tempfile
import time
import unicodedata
import uuid
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta

import sqlalchemy as sa

from . import sinan, sinan_import
from .domain import CSV_COLUMNS, DEPOSIT_COLUMNS, DEPOSIT_TYPES
from .field_store import FieldConflict, FieldStore, FieldStoreError
from .migracao_models import Migracao
from .planning import RG_COLUMNS, read_csv
from .service import PreviewValidationError, preview_csv

MAX_BYTES = 2 * 1024 * 1024
MAX_ROWS = 5000
MAX_OCORRENCIAS = 20000
MAX_GEOCODE = 25
FORMATS = {".csv": "csv", ".txt": "csv", ".xlsx": "xlsx", ".dbf": "dbf"}
DESTINOS = {
    "visitas": "Visitas de campo (LIRAa / tratamento focal)",
    "sinan": "Casos de arboviroses (SINAN)",
    "quarteiroes": "Quarteirões / território",
}


class MigracaoError(ValueError):
    pass


def _key(value):
    text = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^A-Z0-9]+", "_", text.upper()).strip("_")


def _spaces(value):
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _code(value):
    return re.sub(r"[\s/\\]+", "-", _spaces(value).upper()).strip("-")


# ---------------------------------------------------------------- I. Extração

def _cell(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, bytes):
        return value.decode("latin-1")
    return str(value)


def _from_text(content):
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    first = next((line for line in text.splitlines() if line.strip()), "")
    delimiter = max(";,\t|", key=first.count)
    rows = []
    for row in csv.reader(io.StringIO(text, newline=""), delimiter=delimiter):
        rows.append(row)
        if len(rows) > MAX_ROWS + 20:
            break
    return rows, {"codificacao": encoding.replace("-sig", ""), "separador": {"\t": "tabulação"}.get(delimiter, delimiter)}


def _from_xlsx(content):
    from openpyxl import load_workbook
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.worksheets[0]
    rows = []
    for row in sheet.iter_rows(values_only=True):
        rows.append([_cell(value) for value in row])
        if len(rows) > MAX_ROWS + 20:
            break
    workbook.close()
    return rows, {"planilha": sheet.title}


def _from_dbf(content):
    from dbfread import DBF
    with tempfile.NamedTemporaryFile(suffix=".dbf") as tmp:
        tmp.write(content)
        tmp.flush()
        table = DBF(tmp.name, ignore_missing_memofile=True, char_decode_errors="replace")
        if not table.header.language_driver:
            # Exportações antigas sem driver de idioma costumam ser Windows ANSI.
            table = DBF(tmp.name, encoding="cp1252", ignore_missing_memofile=True, char_decode_errors="replace")
        rows = [list(table.field_names)]
        for record in table:
            rows.append([_cell(record[name]) for name in table.field_names])
            if len(rows) > MAX_ROWS + 20:
                break
        return rows, {"codificacao": table.encoding, "registros_dbf": len(rows) - 1}


def extract(content, filename):
    """Converte o arquivo legado em CSV UTF-8 (;) e devolve o texto e o resumo."""
    kind = FORMATS.get(os.path.splitext((filename or "").lower())[1])
    if not kind:
        raise MigracaoError("Envie um arquivo CSV, TXT, XLSX ou DBF exportado do sistema legado.")
    if not content:
        raise MigracaoError("O arquivo está vazio.")
    if len(content) > MAX_BYTES:
        raise MigracaoError("O arquivo excede o limite de 2 MiB por execução. Divida a exportação.")
    try:
        raw, info = {"csv": _from_text, "xlsx": _from_xlsx, "dbf": _from_dbf}[kind](content)
    except MigracaoError:
        raise
    except Exception as exc:
        raise MigracaoError(f"Não foi possível ler o arquivo como {kind.upper()}. Confira a extensão e se a exportação está íntegra.") from exc
    # Planilhas costumam ter título antes do cabeçalho: usa a primeira linha com 2+ colunas preenchidas.
    start = next((i for i, row in enumerate(raw) if sum(bool(_spaces(v)) for v in row) >= 2), None)
    if start is None:
        raise MigracaoError("Não foi encontrado um cabeçalho com colunas no arquivo.")
    header, seen = [], Counter()
    for index, name in enumerate(raw[start]):
        name = _spaces(name) or f"COLUNA_{index + 1}"
        seen[name] += 1
        header.append(name if seen[name] == 1 else f"{name}_{seen[name]}")
    while header and header[-1].startswith("COLUNA_") and not any(len(r) > len(header) - 1 and _spaces(r[len(header) - 1]) for r in raw[start + 1:]):
        header.pop()
    rows, empty = [], 0
    for row in raw[start + 1:]:
        values = [_spaces(v) for v in row[:len(header)]] + [""] * max(0, len(header) - len(row))
        if not any(values):
            empty += 1
            continue
        rows.append(values)
    if not rows:
        raise MigracaoError("O arquivo não contém registros abaixo do cabeçalho.")
    if len(rows) > MAX_ROWS:
        raise MigracaoError(f"O arquivo excede {MAX_ROWS} registros. Divida a exportação.")
    stream = io.StringIO(newline="")
    writer = csv.writer(stream, delimiter=";")
    writer.writerow(header)
    writer.writerows(rows)
    info.update(formato=kind, linhas=len(rows), colunas=header, linhas_antes_cabecalho=start, linhas_vazias=empty)
    return ("﻿" + stream.getvalue()).encode("utf-8"), info


def read_extracted(content):
    reader = csv.reader(io.StringIO(bytes(content).decode("utf-8-sig"), newline=""), delimiter=";")
    header = next(reader)
    return header, list(reader)


# ------------------------------------------------------- II. Mapeamento

_DEPOSIT_SYNONYMS = {code: f"{code.upper()} DEP_{code.upper()} QT_{code.upper()} DEPOSITO_{code.upper()} DEPOSITOS_{code.upper()}" for code, _ in DEPOSIT_TYPES}
TARGETS = {
    "visitas": (
        ("imovel_codigo", "Código do imóvel", True, "COD_IMOVEL CODIMOVEL IMOVEL ID_IMOVEL NU_IMOVEL NUM_IMOVEL CO_IMOVEL CODIGO_IMOVEL SEQ_IMOVEL"),
        ("data_inspecao", "Data da visita", True, "DT_VISITA DATA_VISITA DATA DT_INSPECAO DATA_INSPECAO DT_ATIVIDADE DT_TRABALHO"),
        ("inspecao_codigo", "Código da visita", False, "COD_VISITA ID_VISITA NU_VISITA VISITA COD_INSPECAO INSPECAO ID_INSPECAO NU_BOLETIM"),
        ("origem", "Atividade (LIRAa/focal)", False, "ATIVIDADE TIPO_ATIVIDADE TP_ATIVIDADE ORIGEM TIPO TIPO_VISITA"),
        ("levantamento_codigo", "Levantamento / ciclo", False, "CICLO LEVANTAMENTO COD_LEVANTAMENTO NU_CICLO ID_LEVANTAMENTO LIRAA"),
        ("estrato_codigo", "Estrato", False, "ESTRATO COD_ESTRATO NU_ESTRATO"),
        ("quarteirao", "Quarteirão (enriquecimento)", False, "QUARTEIRAO QUART NU_QUARTEIRAO COD_QUARTEIRAO CO_QUARTEIRAO QUADRA QT"),
        ("bairro", "Bairro", False, "BAIRRO NM_BAIRRO NO_BAIRRO NOME_BAIRRO LOCALIDADE"),
        ("area_codigo", "Área / zona", False, "AREA COD_AREA CO_AREA AREA_CODIGO ZONA SETOR MICROAREA"),
        ("situacao", "Situação do imóvel", False, "SITUACAO PENDENCIA SIT_IMOVEL SITUACAO_IMOVEL STATUS"),
        ("positivo_aedes", "Positivo para Aedes", False, "POSITIVO FOCO AEDES POSITIVO_AEDES RESULTADO FOCO_AEGYPTI AE_AEGYPTI"),
        *((name, f"Recipientes {code.upper()}", False, _DEPOSIT_SYNONYMS[code]) for name, (code, _) in zip(DEPOSIT_COLUMNS, DEPOSIT_TYPES)),
        ("latitude", "Latitude", False, "LATITUDE LAT NU_LATITUDE COORD_Y Y"),
        ("longitude", "Longitude", False, "LONGITUDE LON LNG LONG NU_LONGITUDE COORD_X X"),
        ("logradouro", "Logradouro (geocodificação)", False, "LOGRADOURO ENDERECO NM_LOGRADOURO NM_LOGRADO RUA"),
        ("numero", "Número (geocodificação)", False, "NUMERO NU_NUMERO NUM NU_LOGRADOURO"),
        ("municipio", "Município / IBGE", False, "MUNICIPIO ID_MUNICIP CO_MUNICIPIO COD_IBGE IBGE CD_MUN MUNICIPIO_ID"),
    ),
    "sinan": (
        ("referencia", "Nº da notificação", True, "NU_NOTIFIC NOTIFICACAO NUM_NOTIFICACAO NU_NOTIFICACAO REFERENCIA ID_NOTIFIC"),
        ("doenca", "Agravo / doença", False, "ID_AGRAVO AGRAVO DOENCA CID CID10"),
        ("data_notificacao", "Data da notificação", False, "DT_NOTIFIC DATA_NOTIFICACAO DT_NOTIFICACAO DATA DT_SIN_PRI"),
        ("semana", "Semana epidemiológica", False, "SEM_NOT SEMANA SE SEM_PRI SEMANA_EPIDEMIOLOGICA"),
        ("ano", "Ano", False, "NU_ANO ANO ANO_NOTIFICACAO"),
        ("municipio_ibge", "Município (IBGE)", False, "ID_MUNICIP ID_MN_RESI MUNICIPIO COD_IBGE IBGE MUNICIPIO_IBGE"),
        ("classificacao", "Classificação final", False, "CLASSI_FIN CLASSIFICACAO CLASSIF CLASSIFICACAO_FINAL"),
        ("latitude", "Latitude", False, "LATITUDE LAT NU_LATITUDE COORD_Y Y"),
        ("longitude", "Longitude", False, "LONGITUDE LON LNG LONG NU_LONGITUDE COORD_X X"),
        ("logradouro", "Logradouro (geocodificação)", False, "NM_LOGRADO LOGRADOURO ENDERECO NM_LOGRADOURO RUA"),
        ("numero", "Número (geocodificação)", False, "NU_NUMERO NUMERO NUM"),
        ("bairro", "Bairro (geocodificação)", False, "NM_BAIRRO BAIRRO NO_BAIRRO"),
    ),
    "quarteiroes": (
        ("quarteirao", "Quarteirão", True, "QUARTEIRAO NU_QUARTEIRAO COD_QUARTEIRAO CO_QUARTEIRAO QUART QUADRA QT"),
        ("bairro", "Bairro", True, "BAIRRO NM_BAIRRO NO_BAIRRO NOME_BAIRRO LOCALIDADE"),
        ("area", "Área / zona", True, "AREA COD_AREA CO_AREA ZONA SETOR MICROAREA"),
        ("imoveis", "Quantidade de imóveis", True, "QT_IMOVEIS IMOVEIS NU_IMOVEIS TOTAL_IMOVEIS QTD_IMOVEIS N_IMOVEIS"),
        ("referencia", "Referência", False, "REFERENCIA OBS OBSERVACAO DESCRICAO PONTO_REFERENCIA"),
    ),
}


def suggest_mapping(destino, header):
    keys = {column: _key(column) for column in header}
    mapping, used = {}, set()
    for name, _, _, synonyms in TARGETS[destino]:
        options = [_key(name), *synonyms.split()]
        match = next((column for option in options for column in header
                      if column not in used and keys[column] == option), "")
        mapping[name] = match
        if match:
            used.add(match)
    return mapping


def parse_mapping(destino, header, values):
    mapping = {}
    for name, label, required, _ in TARGETS[destino]:
        column = str(values.get(f"map_{name}", "") or "")
        if column and column not in header:
            raise MigracaoError(f"A coluna escolhida para {label} não existe no arquivo extraído.")
        if required and not column:
            raise MigracaoError(f"Mapeie a coluna obrigatória: {label}.")
        mapping[name] = column
    used = [c for c in mapping.values() if c]
    if len(used) != len(set(used)):
        raise MigracaoError("Cada coluna do arquivo legado pode ser usada em apenas um campo.")
    return mapping


# --------------------------------------------- II. Validação e limpeza

class _Linha:
    def __init__(self, numero, values):
        self.numero, self.values = numero, values
        self.ocorrencias, self.rejeitada = [], False

    def add(self, tipo, campo, original, final, motivo):
        self.ocorrencias.append(dict(linha=self.numero, tipo=tipo, campo=campo, original=str(original or ""),
                                     final=str(final if final is not None else ""), motivo=motivo))

    def padroniza(self, campo, original, final, motivo):
        if str(original) != str(final) and original not in ("", None):
            self.add("PADRONIZADO", campo, original, final, motivo)

    def corrige(self, campo, original, final, motivo):
        self.add("CORRIGIDO", campo, original, final, motivo)

    def enriquece(self, campo, final, motivo):
        self.add("ENRIQUECIDO", campo, "", final, motivo)

    def rejeita(self, campo, original, motivo):
        self.rejeitada = True
        self.add("REJEITADO", campo, original, "", motivo)


def parse_date(value):
    value = _spaces(value)
    match = re.fullmatch(r"(\d{1,2})[/.-](\d{1,2})[/.-](\d{4}|\d{2})", value)
    if match:
        day, month, year = (int(part) for part in match.groups())
        if len(match.group(3)) == 2:
            year += 2000 if year <= 69 else 1900
        return date(year, month, day)
    match = re.fullmatch(r"(\d{4})-(\d{1,2})-(\d{1,2})(?:[ T].*)?", value) or re.fullmatch(r"(\d{4})(\d{2})(\d{2})", value)
    if match:
        return date(*(int(part) for part in match.groups()))
    raise ValueError(value)


def _date(linha, campo, value, required=True):
    if not value:
        if required:
            linha.rejeita(campo, value, "Data obrigatória ausente.")
        return None
    try:
        parsed = parse_date(value)
    except ValueError:
        linha.rejeita(campo, value, "Data inválida (use dd/mm/aaaa).")
        return None
    if not date(2000, 1, 1) <= parsed <= date.today():
        linha.rejeita(campo, value, "Data fora do período aceito (2000 até hoje).")
        return None
    linha.padroniza(campo, value, parsed.isoformat(), "Data convertida para o padrão ISO (aaaa-mm-dd).")
    return parsed


def _number(linha, campo, value):
    text = value.replace(" ", "")
    if "," in text and "." not in text:
        text = text.replace(",", ".")
    number = float(text)
    if not math.isfinite(number):
        raise ValueError(value)
    linha.padroniza(campo, value, text, "Separador decimal padronizado para ponto.")
    return number


def in_brazil(lat, lng):
    return -34 <= lat <= 6 and -74 <= lng <= -28


def _coordinates(linha, lat_raw, lng_raw):
    """Devolve (lat, lng), (None, None) ou False quando a linha foi rejeitada."""
    if not lat_raw and not lng_raw:
        return None, None
    if not lat_raw or not lng_raw:
        linha.rejeita("coordenadas", f"{lat_raw} {lng_raw}".strip(), "Latitude e longitude devem ser informadas juntas.")
        return False
    try:
        lat, lng = _number(linha, "latitude", lat_raw), _number(linha, "longitude", lng_raw)
    except ValueError:
        linha.rejeita("coordenadas", f"{lat_raw}, {lng_raw}", "Coordenada não numérica.")
        return False
    if not in_brazil(lat, lng) and in_brazil(lng, lat):
        linha.corrige("coordenadas", f"{lat}, {lng}", f"{lng}, {lat}", "Latitude e longitude estavam invertidas.")
        lat, lng = lng, lat
    if not in_brazil(lat, lng):
        linha.rejeita("coordenadas", f"{lat_raw}, {lng_raw}", "Coordenadas fora do território brasileiro.")
        return False
    return round(lat, 7), round(lng, 7)


def _enum(value, table):
    key = _key(value)
    return next((name for name, options in table.items() if key in options), None)


ORIGENS = {"LIRAA": {"LIRAA", "LIRA", "LIA", "LI", "LI_T", "LIT", "LEVANTAMENTO", "LEVANTAMENTO_DE_INDICE", "1", "2"},
           "FOCAL": {"FOCAL", "T", "TF", "TRAT", "TRATAMENTO", "TRATAMENTO_FOCAL", "4"}}
SITUACOES = {"INSPECIONADO": {"INSPECIONADO", "INSP", "I", "T", "TRABALHADO", "VISITADO", "V", "NORMAL", "1"},
             "FECHADO": {"FECHADO", "F", "2"}, "RECUSADO": {"RECUSADO", "RECUSA", "R", "3"},
             "NAO_INSPECIONADO": {"NAO_INSPECIONADO", "NI", "NAO_VISITADO"}}
SIM_NAO = {"SIM": {"S", "SIM", "1", "P", "POS", "POSITIVO", "X", "TRUE", "Y"},
           "NAO": {"N", "NAO", "0", "NEG", "NEGATIVO", "FALSE"}}
DOENCAS = {"dengue": {"A90", "A90_0", "A900", "A91", "DENGUE"},
           "chikungunya": {"A92_0", "A920", "CHIKUNGUNYA", "CHIK"},
           "zika": {"A92_8", "A928", "U06_9", "U069", "ZIKA", "ZIKAVIRUS", "ZIKA_VIRUS"}}
CLASSES = {"confirmado": {"1", "10", "11", "12", "13", "CONFIRMADO", "DENGUE", "DENGUE_COM_SINAIS_DE_ALARME", "DENGUE_GRAVE", "CHIKUNGUNYA"},
           "descartado": {"2", "5", "DESCARTADO"},
           "ignorado": {"0", "8", "9", "IGNORADO", "INCONCLUSIVO"},
           "suspeito": {"SUSPEITO", "EM_INVESTIGACAO", "INVESTIGACAO"}}


def epi_week(day):
    """Semana epidemiológica (domingo a sábado; semana 1 tem ao menos 4 dias no ano)."""
    def start(year):
        jan1 = date(year, 1, 1)
        back = (jan1.weekday() + 1) % 7
        return jan1 - timedelta(days=back) if back <= 3 else jan1 + timedelta(days=7 - back)
    year = day.year + (day >= start(day.year + 1)) - (day < start(day.year))
    return f"{year}{(day - start(year)).days // 7 + 1:02d}"


def _municipio(linha, value, ibge, prefeitura, campo):
    """Aceita id interno, IBGE de 7 dígitos ou de 6 (padrão SINAN/DATASUS)."""
    digits = re.sub(r"\D", "", value)
    if not value:
        return True
    if ibge and len(digits) == 6 and ibge.startswith(digits):
        linha.corrige(campo, value, ibge, "Código IBGE de 6 dígitos completado com o dígito verificador.")
        return True
    if digits in {ibge, str(prefeitura)}:
        return True
    linha.rejeita(campo, value, "Registro de outro município; não pertence à prefeitura da migração.")
    return False


def _geocoder_step(linha, geocode, cache, address):
    if not geocode or not address or len(cache) >= MAX_GEOCODE and address not in cache:
        return None
    if address not in cache:
        try:
            result = geocode(address)
        except Exception:
            result = None
        cache[address] = (round(result[0], 7), round(result[1], 7)) if result and in_brazil(*result) else None
    if cache[address]:
        lat, lng = cache[address]
        linha.enriquece("coordenadas", f"{lat:.6f}, {lng:.6f}", "Coordenadas obtidas por geocodificação do endereço.")
    return cache[address]


def _visit(linha, get, mapped, options, ctx, padroes):
    imovel = _code(get("imovel_codigo"))
    linha.padroniza("imovel_codigo", get("imovel_codigo"), imovel, "Código padronizado (maiúsculas, sem espaços).")
    if not imovel:
        linha.rejeita("imovel_codigo", "", "Código do imóvel ausente.")
    elif len(imovel) > 80:
        linha.rejeita("imovel_codigo", imovel, "Código do imóvel com mais de 80 caracteres.")
    visit_date = _date(linha, "data_inspecao", get("data_inspecao"))

    raw = get("origem")
    origem = _enum(raw, ORIGENS) if raw else options["origem"]
    if not raw:
        padroes[f"Atividade padrão da migração: {options['origem']}"] += 1
    elif not origem:
        linha.rejeita("origem", raw, "Atividade desconhecida (use LIRAa ou tratamento focal).")
    elif origem != raw:
        linha.padroniza("origem", raw, origem, "Atividade do legado convertida para o código interno.")

    levantamento = estrato = ""
    if origem == "LIRAA":
        levantamento = _code(get("levantamento_codigo")) or options["levantamento"]
        estrato = _code(get("estrato_codigo")) or options["estrato"]
        linha.padroniza("levantamento_codigo", get("levantamento_codigo"), levantamento, "Código padronizado (maiúsculas, sem espaços).")
        linha.padroniza("estrato_codigo", get("estrato_codigo"), estrato, "Código padronizado (maiúsculas, sem espaços).")
        if not levantamento or not estrato:
            linha.rejeita("levantamento_codigo", "", "Visita LIRAa sem levantamento/estrato e sem valor padrão informado.")
    elif origem == "FOCAL" and (get("levantamento_codigo") or get("estrato_codigo")):
        linha.corrige("levantamento_codigo", f"{get('levantamento_codigo')} {get('estrato_codigo')}".strip(), "",
                      "Tratamento focal não usa levantamento/estrato; valores removidos.")

    raw = get("situacao")
    situacao = _enum(raw, SITUACOES) if raw else ("INSPECIONADO" if not mapped("situacao") else None)
    if not mapped("situacao"):
        padroes["Situação não informada no legado: INSPECIONADO"] += 1
    elif not raw:
        linha.rejeita("situacao", "", "Situação do imóvel ausente.")
    elif not situacao:
        linha.rejeita("situacao", raw, "Situação desconhecida (use trabalhado/inspecionado, fechado, recusado).")
    elif situacao != raw:
        linha.padroniza("situacao", raw, situacao, "Situação do legado convertida para o código interno.")

    deposits = {}
    if any(mapped(name) for name in DEPOSIT_COLUMNS):
        values = {name: get(name) for name in DEPOSIT_COLUMNS}
        if situacao != "INSPECIONADO":
            if any(v not in ("", "0") for v in values.values()):
                linha.corrige("depositos", " ".join(v for v in values.values() if v), "", "Imóvel não inspecionado não registra recipientes; valores removidos.")
        elif any(values.values()):
            for name, value in values.items():
                if not value:
                    linha.corrige(name, "", "0", "Recipiente em branco tratado como zero (demais categorias informadas).")
                    value = "0"
                if not re.fullmatch(r"[0-9]{1,5}", value) or int(value) > 10000:
                    linha.rejeita(name, value, "Quantidade de recipientes inválida (inteiro de 0 a 10.000).")
                else:
                    deposits[name] = int(value)

    raw = get("positivo_aedes")
    positivo = _enum(raw, SIM_NAO) if raw else ""
    if raw and not positivo:
        linha.rejeita("positivo_aedes", raw, "Resultado para Aedes desconhecido (use S/N).")
    if situacao == "INSPECIONADO":
        if not raw and len(deposits) == len(DEPOSIT_COLUMNS):
            positivo = "SIM" if sum(deposits.values()) else "NAO"
            linha.enriquece("positivo_aedes", positivo, "Positividade deduzida das contagens de recipientes.")
        elif not raw:
            linha.rejeita("positivo_aedes", "", "Imóvel inspecionado sem resultado para Aedes.")
        elif positivo != raw:
            linha.padroniza("positivo_aedes", raw, positivo, "Resultado convertido para SIM/NAO.")
        if positivo and len(deposits) == len(DEPOSIT_COLUMNS) and (sum(deposits.values()) > 0) != (positivo == "SIM"):
            linha.rejeita("positivo_aedes", raw, "Resultado para Aedes diverge das contagens de recipientes positivos.")
    elif positivo and situacao:
        linha.corrige("positivo_aedes", raw, "", "Imóvel fechado/recusado não tem resultado; valor removido.")
        positivo = ""

    bairro, area = _spaces(get("bairro")).upper(), _code(get("area_codigo"))
    linha.padroniza("bairro", get("bairro"), bairro, "Texto padronizado (maiúsculas e espaços).")
    linha.padroniza("area_codigo", get("area_codigo"), area, "Código padronizado (maiúsculas, sem espaços).")
    block_code = _code(get("quarteirao"))
    block = ctx["territorio"].get(block_code)
    if block and not bairro:
        bairro = block["bairro"]
        linha.enriquece("bairro", bairro, f"Bairro preenchido pelo quarteirão {block_code} cadastrado.")
    if block and not area:
        area = block["area"]
        linha.enriquece("area_codigo", area, f"Área preenchida pelo quarteirão {block_code} cadastrado.")
    missing = " e quarteirão não cadastrado no território" if block_code and not block else ""
    if not bairro:
        linha.rejeita("bairro", block_code, "Bairro ausente" + missing + ".")
    if not area:
        linha.rejeita("area_codigo", block_code, "Área ausente" + missing + ".")

    if mapped("municipio"):
        _municipio(linha, get("municipio"), ctx["ibge"], ctx["prefeitura"], "municipio")
    coordinates = _coordinates(linha, get("latitude"), get("longitude"))
    if coordinates == (None, None) and not linha.rejeitada:
        address = ", ".join(filter(None, (get("logradouro"), get("numero"), bairro))) if get("logradouro") else ""
        coordinates = _geocoder_step(linha, ctx.get("geocode"), ctx["geocache"], address) or (None, None)

    inspection = _code(get("inspecao_codigo"))
    if not inspection and imovel and visit_date and origem:
        inspection = f"MIG-{origem[0]}-{imovel}-{visit_date:%Y%m%d}"[:80]
        if mapped("inspecao_codigo"):
            linha.enriquece("inspecao_codigo", inspection, "Código da visita gerado a partir de imóvel, atividade e data.")
        else:
            padroes["Código da visita gerado (imóvel + atividade + data)"] += 1
    if linha.rejeitada:
        return None
    data = dict(municipio_id=str(ctx["prefeitura"]), origem=origem, imovel_codigo=imovel, inspecao_codigo=inspection,
                data_inspecao=visit_date.isoformat(), bairro=bairro[:120], area_codigo=area[:80],
                levantamento_codigo=levantamento, estrato_codigo=estrato, situacao=situacao,
                positivo_aedes=positivo, latitude="" if coordinates[0] is None else coordinates[0],
                longitude="" if coordinates[1] is None else coordinates[1])
    if deposits:
        data.update({name: deposits[name] for name in DEPOSIT_COLUMNS})
    return (origem, inspection), data


def _case(linha, get, mapped, options, ctx, padroes):
    raw = get("referencia")
    reference = re.sub(r"[^A-Za-z0-9_-]", "", raw).lstrip("_-")[:64]
    if raw and reference != raw:
        linha.corrige("referencia", raw, reference, "Caracteres inválidos removidos do número da notificação.")
    if not reference:
        linha.rejeita("referencia", raw, "Número da notificação ausente.")

    raw = get("doenca")
    disease = _enum(raw, DOENCAS) if raw else options["doenca"]
    if not raw:
        padroes[f"Agravo padrão da migração: {options['doenca']}"] += 1
    elif not disease:
        linha.rejeita("doenca", raw, "Agravo não é dengue, zika ou chikungunya.")
    elif disease != raw:
        linha.padroniza("doenca", raw, disease, "Agravo (CID/SINAN) convertido para o código interno.")

    if mapped("municipio_ibge") and get("municipio_ibge"):
        _municipio(linha, get("municipio_ibge"), ctx["ibge"], ctx["prefeitura"], "municipio_ibge")
    else:
        padroes[f"Município preenchido pela prefeitura da migração (IBGE {ctx['ibge']})"] += 1

    notified = _date(linha, "data_notificacao", get("data_notificacao"), required=False)
    raw_year = get("ano")
    year = int(raw_year) if re.fullmatch(r"[0-9]{4}", raw_year) else (notified.year if notified else None)
    if raw_year and not re.fullmatch(r"[0-9]{4}", raw_year):
        linha.rejeita("ano", raw_year, "Ano inválido.")
    elif year is None:
        linha.rejeita("ano", "", "Ano ausente e sem data de notificação para deduzi-lo.")
    elif not raw_year:
        linha.enriquece("ano", year, "Ano deduzido da data de notificação.")
    if year and not 2015 <= year <= date.today().year:
        linha.rejeita("ano", year, f"Ano fora do período aceito (2015 a {date.today().year}).")

    raw = get("semana")
    week = re.sub(r"\D", "", raw)
    if week and len(week) <= 2 and year:
        week = f"{year}{int(week):02d}"
        linha.padroniza("semana", raw, week, "Semana epidemiológica padronizada para AAAASS.")
    if not week and notified:
        week = epi_week(notified)
        linha.enriquece("semana", week, "Semana epidemiológica calculada a partir da data de notificação.")
    if week and (len(week) != 6 or not 1 <= int(week[4:]) <= 53 or (year and abs(int(week[:4]) - year) > 1)):
        linha.rejeita("semana", raw, "Semana epidemiológica inválida (use AAAASS).")

    raw = get("classificacao")
    classification = _enum(raw, CLASSES) if raw else "suspeito"
    if not raw:
        linha.corrige("classificacao", "", "suspeito", "Classificação final em branco: caso em investigação (suspeito).")
    elif not classification:
        linha.rejeita("classificacao", raw, "Classificação final desconhecida.")
    elif classification != raw:
        linha.padroniza("classificacao", raw, classification, "Classificação final convertida para o código interno.")

    coordinates = _coordinates(linha, get("latitude"), get("longitude"))
    if coordinates == (None, None) and not linha.rejeitada:
        address = ", ".join(filter(None, (get("logradouro"), get("numero"), get("bairro")))) if get("logradouro") else ""
        coordinates = _geocoder_step(linha, ctx.get("geocode"), ctx["geocache"], address) or (None, None)
        if coordinates == (None, None):
            linha.rejeita("coordenadas", address, "Caso sem coordenadas e sem geocodificação possível do endereço.")
    if linha.rejeitada:
        return None
    return reference, dict(referencia=reference, doenca=disease, municipio_ibge=ctx["ibge"], ano=str(year),
                           semana=week, classificacao=classification, latitude=coordinates[0], longitude=coordinates[1])


def _block(linha, get, mapped, options, ctx, padroes):
    code = _code(get("quarteirao"))
    linha.padroniza("quarteirao", get("quarteirao"), code, "Código padronizado (maiúsculas, sem espaços).")
    bairro = _spaces(get("bairro")).upper()
    linha.padroniza("bairro", get("bairro"), bairro, "Texto padronizado (maiúsculas e espaços).")
    area = _code(get("area"))
    linha.padroniza("area", get("area"), area, "Código padronizado (maiúsculas, sem espaços).")
    for name, value, limit in (("quarteirao", code, 80), ("bairro", bairro, 120), ("area", area, 80)):
        if not value:
            linha.rejeita(name, "", "Campo obrigatório ausente.")
        elif len(value) > limit:
            linha.rejeita(name, value, f"Valor com mais de {limit} caracteres.")
    raw = get("imoveis")
    digits = raw.replace(".", "").replace(" ", "")
    if re.fullmatch(r"[0-9]{1,10}", digits) and 0 < int(digits) <= 2147483647:
        linha.padroniza("imoveis", raw, digits, "Separador de milhar removido.")
    else:
        linha.rejeita("imoveis", raw, "Quantidade de imóveis ausente ou inválida (inteiro maior que zero).")
    reference = _spaces(get("referencia"))[:200]
    if not reference:
        reference = f"Migração {options['fonte']}"[:200]
        padroes["Referência preenchida com o sistema de origem"] += 1
    if linha.rejeitada:
        return None
    return code, dict(quarteirao=code, bairro=bairro, area=area, imoveis=digits, referencia=reference)


ROW_RULES = {"visitas": _visit, "sinan": _case, "quarteiroes": _block}


def _csv_bytes(columns, records, delimiter=";"):
    stream = io.StringIO(newline="")
    writer = csv.DictWriter(stream, fieldnames=columns, delimiter=delimiter, extrasaction="ignore")
    writer.writeheader()
    writer.writerows(records)
    return stream.getvalue().encode("utf-8")


def _columns(destino, records):
    if destino == "visitas":
        return (*CSV_COLUMNS, *DEPOSIT_COLUMNS) if any(DEPOSIT_COLUMNS[0] in r for r in records) else CSV_COLUMNS
    return sinan_import.COLUMNS if destino == "sinan" else RG_COLUMNS


def _reject_accepted(accepted, index, campo, motivo):
    linha, record = accepted.pop(index)
    linha.rejeita(campo, "", motivo)


def _logical_visits(accepted, ctx):
    """III. Validação lógica/física pelo validador da importação LIRAa e pela base já carregada."""
    for _ in range(10):
        records = [record for _, record in accepted]
        if not records:
            return {}
        try:
            report = preview_csv(_csv_bytes(_columns("visitas", records), records), prefeitura_id=ctx["prefeitura"])
        except PreviewValidationError as exc:
            raise MigracaoError(f"O arquivo adequado não passou na validação física: {exc.mensagem}") from exc
        failing = defaultdict(list)
        for error in report["erros"]:
            failing[error["linha"] - 2].append(error)
        for index in sorted(failing, reverse=True):
            errors = failing[index]
            _reject_accepted(accepted, index, errors[0]["campo"], "Validação lógica: " + " ".join(dict.fromkeys(e["mensagem"] for e in errors)))
        if failing:
            continue
        reused = 0
        for index in range(len(report["registros"]) - 1, -1, -1):
            record = report["registros"][index]
            data = {key: record[key] for key in (*CSV_COLUMNS, *DEPOSIT_COLUMNS) if key in record}
            previous = ctx["existentes"].get((data["origem"], data["inspecao_codigo"]))
            if record["origem"] == "LIRAA" and record["levantamento_codigo"] in ctx["planos"]:
                _reject_accepted(accepted, index, "levantamento_codigo", "Levantamento planejado no sistema: importe estas visitas pelo planejamento, com quarteirão.")
            elif previous is not None and previous != data:
                _reject_accepted(accepted, index, "inspecao_codigo", "Visita já carregada anteriormente com dados diferentes.")
            elif previous is not None:
                reused += 1
        return {"ja_carregados": reused, "levantamentos": sorted({r["levantamento_codigo"] for _, r in accepted if r["levantamento_codigo"]})}
    raise MigracaoError("A validação lógica não convergiu. Revise o arquivo de origem.")


def _sinan_groups(records):
    groups = defaultdict(list)
    for record in records:
        groups[(record["doenca"], record["ano"])].append(record)
    chunks = []
    for (disease, year), items in sorted(groups.items()):
        for start in range(0, len(items), sinan_import.MAX_CASES):
            chunks.append(((disease, year), items[start:start + sinan_import.MAX_CASES]))
    return chunks


def _logical_cases(accepted, ctx):
    for _ in range(len(accepted) + 1):
        failed = None
        for (disease, year), items in _sinan_groups([r for _, r in accepted]):
            try:
                sinan_import.parse_cases(_csv_bytes(sinan_import.COLUMNS, items), {"disease": disease, "ibge": ctx["ibge"], "year": int(year)}, ctx["demo"])
            except ValueError as exc:
                match = re.match(r"Linha (\d+):", str(exc))
                if not match:
                    raise MigracaoError(f"Os casos adequados não passaram na validação do SINAN: {exc}") from exc
                target = items[int(match.group(1)) - 2]
                failed = (next(i for i, (_, r) in enumerate(accepted) if r is target), str(exc).split(": ", 1)[-1])
                break
        if not failed:
            return {"grupos": [{"doenca": d, "ano": y, "casos": len(items)} for (d, y), items in _sinan_groups([r for _, r in accepted])]}
        _reject_accepted(accepted, failed[0], "referencia", "Validação SINAN: " + failed[1])
    raise MigracaoError("A validação dos casos não convergiu.")


def _logical_blocks(accepted, ctx):
    if accepted:
        records = [r for _, r in accepted]
        try:
            read_csv(_csv_bytes(RG_COLUMNS, records), RG_COLUMNS)
        except FieldConflict as exc:
            raise MigracaoError(f"O território adequado não passou na validação: {exc}") from exc
    return {"ja_carregados": sum(r["quarteirao"] in ctx["existentes"] for _, r in accepted)}


LOGICAL = {"visitas": _logical_visits, "sinan": _logical_cases, "quarteiroes": _logical_blocks}


def validate(destino, extraido, mapping, options, ctx):
    """Etapas II e III sem gravar: devolve (relatório, CSV no layout interno)."""
    header, rows = read_extracted(extraido)
    index = {name: header.index(column) for name, column in mapping.items() if column}
    ctx = {**ctx, "geocache": {}}
    padroes, lines, accepted, seen = Counter(), [], [], {}
    for number, values in enumerate(rows, 2):
        linha = _Linha(number, values)
        lines.append(linha)

        def get(name, values=values):
            return _spaces(values[index[name]]) if name in index and index[name] < len(values) else ""

        result = ROW_RULES[destino](linha, get, lambda name: name in index, options, ctx, padroes)
        if result is None:
            continue
        key, record = result
        if key in seen:
            first_line, first_record = seen[key]
            if first_record == record:
                linha.add("DUPLICADO", "registro", "", "", f"Duplicado da linha {first_line}; descartado.")
            else:
                linha.rejeita("registro", "", f"Mesmo identificador da linha {first_line} com dados diferentes.")
            continue
        seen[key] = (number, record)
        accepted.append((linha, record))
    logical = LOGICAL[destino](accepted, ctx)
    records = [record for _, record in accepted]
    columns = list(_columns(destino, records))
    ocorrencias = [item for linha in lines for item in linha.ocorrencias]
    kinds = lambda kind: sum(any(o["tipo"] == kind for o in linha.ocorrencias) for linha in lines if not linha.rejeitada)
    report = {
        "destino": destino,
        "lidos": len(rows),
        "validos": len(records),
        "corrigidos": kinds("CORRIGIDO"),
        "padronizados": kinds("PADRONIZADO"),
        "enriquecidos": kinds("ENRIQUECIDO"),
        "duplicados": sum(any(o["tipo"] == "DUPLICADO" for o in linha.ocorrencias) for linha in lines),
        "rejeitados": sum(linha.rejeitada for linha in lines),
        "sem_coordenadas": sum(r.get("latitude") in ("", None) for r in records) if destino != "quarteiroes" else None,
        "padroes": dict(padroes),
        "regras": [[tipo, motivo, count] for (tipo, motivo), count in Counter(
            (o["tipo"], re.sub(r"\d+", "N", o["motivo"]) if o["tipo"] in {"DUPLICADO", "ENRIQUECIDO"} else o["motivo"])
            for o in ocorrencias).most_common()],
        "ocorrencias": ocorrencias[:MAX_OCORRENCIAS],
        "ocorrencias_truncadas": len(ocorrencias) > MAX_OCORRENCIAS,
        "colunas": columns,
        "amostra": [[record.get(c, "") for c in columns] for record in records[:10]],
        "opcoes": {k: v for k, v in options.items() if k != "fonte"},
        **logical,
    }
    return report, _csv_bytes(columns, records)


def report_csv(report):
    stream = io.StringIO(newline="")
    writer = csv.writer(stream, delimiter=";")

    def safe(value):
        value = str(value)
        return "'" + value if value.startswith(("=", "+", "-", "@")) and not re.fullmatch(r"-?[0-9.,]+", value) else value

    writer.writerow(["resumo", "lidos", report["lidos"], "validos", report["validos"], "corrigidos", report["corrigidos"],
                     "duplicados", report["duplicados"], "rejeitados", report["rejeitados"], "enriquecidos", report["enriquecidos"]])
    writer.writerow(["linha", "tipo", "campo", "valor_original", "valor_final", "motivo"])
    for item in report["ocorrencias"]:
        writer.writerow([item["linha"], item["tipo"], item["campo"], safe(item["original"]), safe(item["final"]), safe(item["motivo"])])
    return ("﻿" + stream.getvalue()).encode("utf-8")


# --------------------------------------------------------- Persistência

class MigracaoStore(FieldStore):
    PUBLIC = ("id", "prefeitura", "demo", "actor", "destino", "fonte", "filename", "formato", "sha256", "etapa",
              "extracao", "mapeamento", "relatorio", "resultado", "criado", "atualizado")

    def create(self, *, prefeitura, demo, actor, destino, fonte, filename, content):
        if destino not in DESTINOS:
            raise MigracaoError("Selecione um destino válido.")
        fonte = _spaces(fonte)
        if not fonte or len(fonte) > 120:
            raise MigracaoError("Informe o sistema legado de origem (até 120 caracteres).")
        extraido, info = extract(content, filename)
        table, now = Migracao.__table__, time.time()
        identifier = uuid.uuid4().hex
        with self.connection(write_prefeitura=prefeitura) as conn:
            if conn.scalar(sa.select(sa.func.count()).select_from(table).where(table.c.prefeitura == prefeitura)) >= 500:
                raise MigracaoError("Limite de 500 execuções de migração por prefeitura. Solicite manutenção.")
            conn.execute(table.insert().values(
                id=identifier, prefeitura=prefeitura, demo=bool(demo), actor=str(actor), destino=destino, fonte=fonte,
                filename=(filename or "arquivo")[:120], formato=info["formato"], sha256=hashlib.sha256(content).hexdigest(),
                original=content, extraido=extraido, etapa="EXTRAIDO", extracao=info,
                mapeamento={"colunas": suggest_mapping(destino, info["colunas"])}, criado=now, atualizado=now))
        return identifier

    def get(self, identifier, prefeitura):
        table = Migracao.__table__
        with self.connection() as conn:
            row = conn.execute(sa.select(*(table.c[n] for n in self.PUBLIC), table.c.extraido, table.c.adequado)
                               .where(table.c.id == identifier, table.c.prefeitura == prefeitura)).mappings().first()
        return dict(row) if row else None

    def file(self, identifier, prefeitura, column):
        table = Migracao.__table__
        with self.connection() as conn:
            value = conn.scalar(sa.select(table.c[column]).where(table.c.id == identifier, table.c.prefeitura == prefeitura))
        return bytes(value) if value is not None else None

    def save(self, identifier, prefeitura, **values):
        table = Migracao.__table__
        with self.connection(write_prefeitura=prefeitura) as conn:
            conn.execute(table.update().where(table.c.id == identifier, table.c.prefeitura == prefeitura)
                         .values(atualizado=time.time(), **values))

    def history(self, prefeitura):
        table = Migracao.__table__
        with self.connection() as conn:
            rows = conn.execute(sa.select(table.c.id, table.c.destino, table.c.fonte, table.c.filename, table.c.formato,
                                          table.c.etapa, table.c.demo, table.c.criado, table.c.relatorio, table.c.resultado)
                                .where(table.c.prefeitura == prefeitura).order_by(table.c.criado.desc()).limit(20)).mappings()
            return [dict(row, relatorio={k: (row["relatorio"] or {}).get(k) for k in ("lidos", "validos", "rejeitados")})
                    for row in rows]

    def context(self, destino, prefeitura, demo):
        """Dados já cadastrados usados no enriquecimento e na validação lógica."""
        from .planning_models import CampoPlano, CampoQuarteirao
        blocks = CampoQuarteirao.__table__
        with self.connection() as conn:
            territory = {row.code: {"bairro": row.bairro.upper(), "area": row.area}
                         for row in conn.execute(sa.select(blocks.c.code, blocks.c.bairro, blocks.c.area)
                                                 .where(blocks.c.prefeitura == prefeitura, blocks.c.demo == demo))}
            plans = set(conn.scalars(sa.select(CampoPlano.code).where(CampoPlano.prefeitura == prefeitura, CampoPlano.demo == demo)))
        context = {"territorio": territory, "planos": plans, "existentes": {}}
        if destino == "visitas":
            context["existentes"] = {(r["origem"], r["inspecao_codigo"]): r for r in self.records(prefeitura, demo=demo)}
        elif destino == "quarteiroes":
            context["existentes"] = set(territory)
        return context


def load(migracao, *, actor, engine=None, sinan_store=None, owner=None):
    """III. Carga do CSV adequado pelos fluxos oficiais de cada destino."""
    destino, prefeitura, demo = migracao["destino"], migracao["prefeitura"], migracao["demo"]
    content = bytes(migracao["adequado"] or b"")
    report = migracao["relatorio"] or {}
    if not report.get("validos"):
        raise MigracaoError("Não há registros válidos para carregar.")
    if destino == "visitas":
        store = FieldStore(engine)
        batch = store.prepare(content, prefeitura=prefeitura, actor=actor, source=f"Migração {migracao['id'][:8]} · {migracao['fonte']}"[:120],
                              filename=f"migracao-{migracao['id'][:12]}.csv", demo=demo)
        if batch["status"] == "REJEITADO":
            first = (batch["report"].get("erros") or [{}])[0]
            raise MigracaoError("A carga foi recusada pela validação de campo: " + first.get("mensagem", "revise o arquivo."))
        done = store.confirm(batch["id"], prefeitura=prefeitura, actor=actor)
        return {"inseridos": done["inserted"], "reaproveitados": done["reused"], "lote": done["id"],
                "levantamentos": report.get("levantamentos", [])}
    if destino == "quarteiroes":
        from .planning import PlanningStore
        PlanningStore(engine).mutate(prefeitura, demo, actor, "importar_territorio", {}, content)
        return {"inseridos": report["validos"] - report.get("ja_carregados", 0), "reaproveitados": report.get("ja_carregados", 0)}
    rows = list(csv.DictReader(io.StringIO(content.decode("utf-8")), delimiter=";"))
    inserted = reused = 0
    jobs = []
    for (disease, year), items in _sinan_groups(rows):
        chunk = _csv_bytes(sinan_import.COLUMNS, items)
        scope = {"disease": disease, "ibge": items[0]["municipio_ibge"], "year": int(year)}
        job_id, existing = sinan_store.create(owner, kind="csv", **scope, demo=demo, fingerprint=hashlib.sha256(chunk).hexdigest())
        if not existing:
            sinan_import.run_cases(sinan_store, job_id, scope, chunk, demo)
        job = sinan_store.get(job_id, owner)
        if job["status"] != "completed":
            raise MigracaoError(f"O SINAN recusou o lote {disease}/{year}: {job['message']}")
        reused, inserted = reused + len(items) * existing, inserted + len(items) * (not existing)
        jobs.append({"id": job_id, "doenca": disease, "ano": int(year), "casos": len(items), "existente": existing})
    return {"inseridos": inserted, "reaproveitados": reused, "lotes_sinan": jobs, "ibge": rows[0]["municipio_ibge"]}


# ----------------------------------------------- Exemplos do sistema legado

def _dbf(fields, rows):
    """dBase III mínimo (campos caractere, Windows ANSI), como os exportados pelo SISPNCD."""
    encoded = [[str(value).encode("cp1252") for value in row] for row in rows]
    widths = [max([1] + [len(row[i]) for row in encoded]) for i in range(len(fields))]
    today = date.today()
    out = bytearray(struct.pack("<BBBBIHH20x", 3, today.year - 1900, today.month, today.day, len(rows),
                                32 + 32 * len(fields) + 1, 1 + sum(widths)))
    out[29] = 0x03
    for name, width in zip(fields, widths):
        out += struct.pack("<11sc4xBB14x", name.encode("ascii"), b"C", width, 0)
    out += b"\r"
    for row in encoded:
        out += b" " + b"".join(value.ljust(width) for value, width in zip(row, widths))
    return bytes(out + b"\x1a")


def sample_file(destino, ibge=None):
    """Exportações fictícias com erros propositais para demonstrar as três etapas."""
    lat, lng = (-8.3558, -36.6978)
    if destino == "visitas":
        fields = ["COD_IMOVEL", "DT_VISITA", "ATIVIDADE", "CICLO", "ESTRATO", "QUARTEIRAO", "BAIRRO", "ZONA", "LOGRADOURO",
                  "NUMERO", "SITUACAO", "FOCO", "A1", "A2", "B", "C", "D1", "D2", "E", "LATITUDE", "LONGITUDE"]
        z = ["0"] * 7
        base = ["LIRAA 10/2026", "E1"]
        rows = [
            ["000101", "05/10/2026", "LI", *base, "101", "Centro", "Z-01", "Rua Duque de Caxias", "10", "T", "N", *z, f"{lat - .0003:.4f}", f"{lng - .0003:.4f}"],
            ["000102", "05/10/2026", "LI", *base, "101", "Centro", "Z-01", "Rua Duque de Caxias", "52", "T", "S", "0", "1", "1", "0", "0", "0", "0", "-8,3563", "-36,6985"],
            ["000103", "05/10/2026", "LI", *base, "101", "centro ", "z-01", "Rua Duque de Caxias", "80", "F", "N", "", "", "", "", "", "", "", "", ""],
            ["000104", "06/10/2026", "LI", *base, "102", "", "", "Rua Barão de Vila Bela", "120", "T", "N", *z, "", ""],
            ["000102", "05/10/2026", "LI", *base, "101", "Centro", "Z-01", "Rua Duque de Caxias", "52", "T", "S", "0", "1", "1", "0", "0", "0", "0", "-8,3563", "-36,6985"],
            ["000106", "31/02/2026", "LI", *base, "102", "São Sebastião", "Z-01", "", "", "T", "N", *z, "", ""],
            ["", "06/10/2026", "LI", *base, "102", "São Sebastião", "Z-01", "", "", "T", "N", *z, "", ""],
            ["000108", "07/10/2026", "LI", *base, "102", "São Sebastião", "Z-01", "", "", "T", "", "0", "0", "0", "0", "0", "2", "0", f"{lat + .0012:.4f}", f"{lng + .0010:.4f}"],
            ["000109", "7/10/26", "LI", *base, "103", "Prado", "Z-02", "", "", "T", "N", *z, f"{lng + .0021:.4f}", f"{lat + .0018:.4f}"],
            ["000110", "08/10/2026", "TRAT", *base, "103", "Prado", "Z-02", "", "", "T", "S", "0", "0", "1", "0", "0", "0", "0", f"{lat + .0025:.4f}", f"{lng + .0019:.4f}"],
            ["000111", "08/10/2026", "LI", *base, "103", "Prado", "Z-02", "", "", "X", "N", *z, "", ""],
            ["000112", "08/10/2026", "LI", *base, "103", "Prado", "Z-02", "", "", "T", "N", *z, "48.8566", "2.3522"],
            ["000113", "09/10/2026", "LI", *base, "103", "Prado", "Z-02", "", "", "R", "", "", "", "", "", "", "", "", "", ""],
            ["000114", "09/10/2026", "LI", *base, "999", "", "", "", "", "T", "N", *z, "", ""],
        ]
        return _dbf(fields, rows), "sispncd-visitas-EXEMPLO.dbf", "application/x-dbf"
    if destino == "sinan":
        stream = io.StringIO(newline="")
        writer = csv.writer(stream, delimiter=";")
        writer.writerow(["NU_NOTIFIC", "ID_AGRAVO", "DT_NOTIFIC", "SEM_NOT", "NU_ANO", "ID_MUNICIP", "CLASSI_FIN",
                         "NM_BAIRRO", "NM_LOGRADO", "NU_NUMERO", "LATITUDE", "LONGITUDE"])
        ibge6 = (ibge or "2610905")[:6]
        point = lambda dy, dx: (f"{lat + dy:.4f}".replace(".", ","), f"{lng + dx:.4f}".replace(".", ","))
        writer.writerows([
            ["6543201", "A90", "05/10/2026", "202640", "2026", ibge6, "10", "Centro", "Rua Duque de Caxias", "10", *point(-.0004, -.0002)],
            ["6543202", "A90", "06/10/2026", "", "2026", ibge6, "", "Centro", "Rua Duque de Caxias", "200", *point(.0006, .0004)],
            ["6543203", "A92.0", "07/10/2026", "202640", "2026", ibge6, "5", "Prado", "Rua do Prado", "15", *point(.0019, .0021)],
            ["6543201", "A90", "05/10/2026", "202640", "2026", ibge6, "10", "Centro", "Rua Duque de Caxias", "10", *point(-.0004, -.0002)],
            ["6543205", "A90", "32/13/2026", "", "2026", ibge6, "10", "Centro", "Rua Duque de Caxias", "30", *point(.0001, .0001)],
            ["6543206", "A90", "08/10/2026", "202641", "2026", ibge6, "1", "São Sebastião", "Rua Barão de Vila Bela", "120", "", ""],
            ["6543207", "A90", "08/10/2026", "202641", "2026", "260120", "1", "Centro", "Rua Cel. Antônio Japiassu", "5", *point(.0003, .0003)],
            ["6543208", "A90", "09/10/2026", "202641", "2026", ibge6, "10", "Centro", "Rua Duque de Caxias", "44", "48,8566", "2,3522"],
            ["6543 209", "A90", "09/10/2026", "202641", "2026", ibge6, "10", "Prado", "Rua do Prado", "90", *point(.0022, .0018)],
            ["6543210", "A90", "09/10/2026", "202641", "2026", ibge6, "99", "Prado", "Rua do Prado", "91", *point(.0023, .0017)],
        ])
        return stream.getvalue().encode("cp1252"), "sinan-notificacoes-EXEMPLO.csv", "text/csv"
    from openpyxl import Workbook
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Quarteiroes"
    sheet.append(["Cadastro de quarteirões - SISLOC / SMS (exemplo fictício)"])
    sheet.append([])
    sheet.append(["NU_QUARTEIRAO", "BAIRRO", "ZONA", "QT_IMOVEIS", "REFERENCIA"])
    for row in ([101, "Centro", "Z-01", 120, "Praça Dom José Lopes"], [102, "São Sebastião", "Z-01", 85, "Igreja de São Sebastião"],
                [103, "  prado ", "z-02", "1.050", None], [101, "Centro", "Z-01", 120, "Praça Dom José Lopes"],
                [104, None, "Z-02", 40, "Sem bairro no legado"], [105, "Cohab", "Z-03", 0, "Quantidade zerada"],
                [106, "Cohab", "Z-03", "abc", "Quantidade não numérica"]):
        sheet.append(row)
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue(), "sisloc-quarteiroes-EXEMPLO.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
