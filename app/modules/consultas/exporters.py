from io import BytesIO
from xml.sax.saxutils import escape
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A3, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle
from app.shared.pdf_watermark import municipal_canvas, report_prefeitura_id


def excel(result):
    book = Workbook(); sheet = book.active; sheet.title = 'Consulta'
    def safe(value):
        value = str(value)
        return "'" + value if value.lstrip().startswith(('=', '+', '-', '@')) else value
    sheet.append([c['label'] for c in result['columns']])
    for row in result['rows']: sheet.append([safe(v) for v in row])
    sheet.freeze_panes = 'A2'; sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='174967')
        sheet.column_dimensions[cell.column_letter].width = 24
    groups = book.create_sheet('Quebras e totais')
    groups.append(['Grupo', 'Registros', 'Campo', 'Soma'])
    for g in result['groups']:
        for field, total in g['totals'].items(): groups.append([safe(g['label']), g['count'], field, total])
        if not g['totals']: groups.append([safe(g['label']), g['count']])
    meta = book.create_sheet('Critérios')
    import json
    meta.append(['Gerado em (UTC)', result['generated_at']]); meta.append(['Registros', result['total']])
    meta.append(['Definição', json.dumps(result['spec'], ensure_ascii=False)])
    stream = BytesIO(); book.save(stream); stream.seek(0); return stream


def pdf(result, user):
    from .service import QueryError
    if result['total'] > 2000:
        raise QueryError('O PDF e o e-mail aceitam até 2.000 registros por geração. Restrinja os filtros ou use Excel.')
    stream = BytesIO(); styles = getSampleStyleSheet()
    styles['BodyText'].fontSize = 8; styles['BodyText'].leading = 10
    def p(value): return Paragraph(escape(str(value)), styles['BodyText'])
    story = [Paragraph('Consulta personalizada', styles['Title']),
        p(f"{result['total']} registros · Gerado em {result['generated_at']}"), Spacer(1, 10)]
    criteria = ', '.join(f'{k}: {v}' for k,v in result['spec']['filters'].items() if v)
    if criteria: story.extend([p(criteria), Spacer(1, 8)])
    rows = [[p(c['label']) for c in result['columns']]]
    group_starts = {g['start']: g for g in result['groups']} if result['spec']['group'] else {}
    for i,row in enumerate(result['rows']):
        if i in group_starts:
            g = group_starts[i]
            rows.append([p(f"{g['label']} ({g['count']} registros)")] + ['']*(len(row)-1))
        rows.append([p(v) for v in row])
    table = LongTable(rows, repeatRows=1, colWidths=[(landscape(A3)[0]-60)/len(result['columns'])]*len(result['columns']))
    table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#def2f5')),
        ('GRID',(0,0),(-1,-1),0.3,colors.HexColor('#cad6dd')),('VALIGN',(0,0),(-1,-1),'TOP')]))
    story.append(table)
    for g in result['groups']:
        totals = '; '.join(f'{key}: {value:g}' for key,value in g['totals'].items())
        if totals: story.extend([Spacer(1, 8), p(f"Somas — {g['label']}: {totals}")])
    SimpleDocTemplate(stream, pagesize=landscape(A3), leftMargin=30, rightMargin=30, topMargin=30, bottomMargin=30).build(
        story, canvasmaker=municipal_canvas(report_prefeitura_id(result['prefeitura_ids'], user), user=user))
    stream.seek(0); return stream
